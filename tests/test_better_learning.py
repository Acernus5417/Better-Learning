"""Observable I/O and coverage regressions, not a test of teaching quality."""
from __future__ import annotations

import argparse
import importlib.util
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[1] / 'better-learning' / 'scripts'
sys.path.insert(0, str(SCRIPTS))
from _common import atomic_text, extraction_path, read_json, write_json
from assemble_knowledge import build_content
from extract_materials import extract
from inventory_materials import inventory as actual_inventory

def inventory(course, inputs):
    result = actual_inventory(course, inputs)
    (course / "_工作区/课程配置.json").unlink(missing_ok=True)
    return result

from validate_package import REQUIRED, validate
from transcribe_materials import initialize as init_transcripts, record as record_transcript


def options(start=1, end=None, max_chars=256):
    return argparse.Namespace(start=start, end=end, max_chars=max_chars, scale=1,
                              ocr='never', ocr_lang='eng')


class WorkflowTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='better-learning-test-')
        self.root = Path(self.temp.name)
        self.inputs = self.root / 'inputs'
        self.inputs.mkdir()
        self.course = self.root / 'course'

    def tearDown(self):
        self.temp.cleanup()

    def register(self, name='notes.md', text='A rate compares a change to the interval over which it occurs.'):
        source = self.inputs / name
        source.write_text(text, encoding='utf-8')
        inventory(self.course, [self.inputs])
        return read_json(self.course / '_工作区' / '资料索引.json')['sources'][0]

    def mapping(self, source):
        return read_json(extraction_path(self.course, source) / '定位映射.json')

    def jsonl(self, name, rows):
        atomic_text(self.course / '_工作区' / name,
                    ''.join(json.dumps(row, ensure_ascii=False) + '\n' for row in rows))

    def valid_package(self):
        source = self.register()
        extract(self.course, source['id'], options())
        init_transcripts(self.course)
        atomic_text(self.course / '_工作区/transcript-input.md', Path(source['path']).read_text(encoding='utf-8'))
        record_transcript(self.course, source['id'], 'U00001', '_工作区/transcript-input.md')
        for name in REQUIRED:
            atomic_text(self.course / name, '# ' + name + '\n\n结构测试样本，不是教学质量验证。\n')
        atomic_text(self.course / '_工作区/章节知识/KC-001.md',
                    '# Rates\n\n<a id="K-000001"></a>\nA rate compares a change to an interval.\n')
        atomic_text(self.course / '学习文档/01-Rates.md',
                    '# Rates\n\n<a id="K-000001"></a>\nRate example.\n\n[Card](../核心知识点.md#KP-01-01-01)\n')
        atomic_text(self.course / '核心知识点.md',
                    '# Cards\n\n<a id="KP-01-01-01"></a>\n## Rate\n[Lesson](学习文档/01-Rates.md#K-000001)\n')
        index = {'schema_version': 1,
                 'chapters': [{'id': 'KC-001', 'title': 'Rates', 'order': 1,
                               'path': '_工作区/章节知识/KC-001.md', 'status': 'complete'}],
                 'lessons': [{'id': 'LC-001', 'title': 'Rates', 'order': 1,
                              'path': '学习文档/01-Rates.md', 'status': 'complete'}]}
        write_json(self.course / '_工作区/章节索引.json', index)
        knowledge = {'id': 'K-000001', 'title': 'Rate', 'chapter_id': 'KC-001',
                     'anchor': 'K-000001', 'kind': 'material', 'status': 'verified',
                     'source_refs': [{'source_id': source['id'], 'unit_id': 'U00001'}],
                     'prerequisites': [], 'lesson_ids': ['LC-001'], 'core': True,
                     'card_id': 'KP-01-01-01', 'card_anchor': 'KP-01-01-01', 'core_reason': 'Course goal'}
        self.jsonl('知识索引.jsonl', [knowledge])
        self.jsonl('覆盖台账.jsonl', [{'source_id': source['id'], 'source_hash': source['sha256'],
              'unit_id': 'U00001', 'status': 'covered', 'knowledge_ids': ['K-000001'],
              'reason': '', 'visual_reviewed': False, 'evidence': 'Read all source text in this test fixture.'}])
        write_json(self.course / '_工作区/生成进度.json', {'schema_version': 1,
                   'confirmations': {'goal': True, 'materials': True, 'baseline': True},
                   'stage': 'validation', 'blockers': []})
        atomic_text(self.course / '知识内容.md', build_content(self.course))
        return source, knowledge

    def test_inventory_stable_ids_and_invalidation(self):
        source = self.register()
        state_path = self.course / '_工作区/生成进度.json'
        state = read_json(state_path)
        state['confirmations'] = {'goal': True, 'materials': True, 'baseline': True}
        state['learner_note'] = 'preserve me'
        write_json(state_path, state)
        result = inventory(self.course, [self.inputs])
        self.assertFalse(result['changed'])
        self.assertTrue(read_json(state_path)['confirmations']['materials'])
        Path(source['path']).write_text('Changed definition.', encoding='utf-8')
        result = inventory(self.course, [self.inputs])
        updated = read_json(self.course / '_工作区/资料索引.json')['sources'][0]
        self.assertTrue(result['changed'])
        self.assertEqual(source['id'], updated['id'])
        self.assertNotEqual(source['sha256'], updated['sha256'])
        state = read_json(state_path)
        self.assertFalse(state['confirmations']['materials'])
        self.assertEqual(state['learner_note'], 'preserve me')

    def test_text_batches_resume_without_losing_any_character(self):
        original = '这是包含连续公式和长题干的材料。\n' * 300
        source = self.register(text=original)
        result = extract(self.course, source['id'], options(end=2))
        self.assertEqual(result['next_start'], 3)
        mapping = self.mapping(source)
        self.assertEqual(mapping['units'][2]['status'], 'pending')
        extract(self.course, source['id'], options(start=3, end=mapping['total_units']))
        mapping = self.mapping(source)
        reconstructed = ''.join((self.course / chunk).read_text(encoding='utf-8')
                                for unit in mapping['units'] for chunk in unit['chunks'])
        self.assertEqual(original, reconstructed)
        self.assertTrue(all(u['status'] == 'extracted' for u in mapping['units']))
        self.assertFalse((self.course / '_工作区/覆盖台账.jsonl').exists())

    def test_changed_input_refuses_stale_extraction(self):
        source = self.register()
        Path(source['path']).write_text('Changed', encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'changed'):
            extract(self.course, source['id'], options())

    def test_unsupported_format_remains_blocked(self):
        source = self.register(name='old.ppt', text='not a readable presentation')
        result = extract(self.course, source['id'], options())
        self.assertEqual(result['processed'][0]['status'], 'blocked')
        self.assertTrue(self.mapping(source)['units'][0]['warnings'])

    @unittest.skipUnless(all(importlib.util.find_spec(m) for m in ['pypdf', 'pypdfium2', 'reportlab', 'PIL']),
                         'PDF fixture requires bundled optional PDF dependencies')
    def test_mixed_pdf_renders_both_pages_and_scan_is_not_silently_empty(self):
        from PIL import Image, ImageDraw
        from reportlab.pdfgen.canvas import Canvas
        raster = self.inputs / 'page.png'
        image = Image.new('RGB', (640, 240), 'white')
        ImageDraw.Draw(image).text((20, 50), 'Scanned formula: speed = distance / time', fill='black')
        image.save(raster)
        pdf = self.inputs / 'mixed.pdf'
        canvas = Canvas(str(pdf))
        canvas.drawString(40, 700, 'Text layer: distance traveled per unit of time is speed.')
        canvas.showPage()
        canvas.drawImage(str(raster), 40, 500, width=480, height=180)
        canvas.save()
        inventory(self.course, [pdf])
        source = read_json(self.course / '_工作区/资料索引.json')['sources'][0]
        result = extract(self.course, source['id'], options(end=2))
        self.assertEqual(result['total_units'], 2)
        mapping = self.mapping(source)
        for unit in mapping['units']:
            self.assertTrue(unit['visual_review_required'])
            self.assertEqual(unit['status'], 'needs_review')
            with Image.open(self.course / unit['images'][0]) as rendered:
                self.assertGreater(rendered.width, 100)
        scan = mapping['units'][1]
        self.assertTrue(scan['warnings'])
        self.assertEqual((self.course / scan['chunks'][0]).read_text(encoding='utf-8'), '')

    @unittest.skipUnless(importlib.util.find_spec('pptx'), 'PPTX fixture requires python-pptx')
    def test_pptx_text_and_notes(self):
        from pptx import Presentation
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = 'Rate'
        slide.placeholders[1].text = 'Change divided by elapsed time'
        slide.notes_slide.notes_text_frame.text = 'Teacher note: check the time unit.'
        path = self.inputs / 'deck.pptx'
        deck.save(path)
        inventory(self.course, [path])
        source = read_json(self.course / '_工作区/资料索引.json')['sources'][0]
        extract(self.course, source['id'], options())
        units = self.mapping(source)['units']
        unit = units[0]
        text = ''.join((self.course / c).read_text(encoding='utf-8') for u in units for c in u['chunks'])
        self.assertIn('Change divided', text)
        self.assertIn('Teacher note', text)
        self.assertEqual(unit['kind'], 'text')
        self.assertFalse(unit['visual_review_required'])

    def test_epub_spine_order(self):
        path = self.inputs / 'book.epub'
        with zipfile.ZipFile(path, 'w') as z:
            z.writestr('META-INF/container.xml', '<container><rootfiles><rootfile full-path="OPS/book.opf"/></rootfiles></container>')
            z.writestr('OPS/book.opf', '<package><manifest><item id="a" href="a.xhtml" media-type="application/xhtml+xml"/><item id="b" href="b.xhtml" media-type="application/xhtml+xml"/><item id="notes" href="notes.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="b"/><itemref idref="a"/></spine></package>')
            z.writestr('OPS/a.xhtml', '<html><body><p>Second chapter</p></body></html>')
            z.writestr('OPS/b.xhtml', '<html><body><p>First chapter</p></body></html>')
            z.writestr('OPS/notes.xhtml', '<html><body><p>Essential footnote outside the spine</p></body></html>')
        inventory(self.course, [path])
        source = read_json(self.course / '_工作区/资料索引.json')['sources'][0]
        extract(self.course, source['id'], options(end=2))
        mapping = self.mapping(source)
        first = (self.course / mapping['units'][0]['chunks'][0]).read_text(encoding='utf-8')
        self.assertIn('First chapter', first)
        self.assertEqual(mapping['total_units'], 3)
        self.assertEqual(mapping['units'][2]['status'], 'pending')
        extract(self.course, source['id'], options(start=3))
        notes_unit = self.mapping(source)['units'][2]
        self.assertIn('Essential footnote', (self.course / notes_unit['chunks'][0]).read_text(encoding='utf-8'))

    @unittest.skipUnless(importlib.util.find_spec('docx'), 'DOCX fixture requires python-docx')
    def test_docx_paragraphs_and_table(self):
        from docx import Document
        document = Document()
        document.add_paragraph('Definition: speed is distance per unit time.')
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Distance'
        table.cell(0, 1).text = 'Time'
        path = self.inputs / 'notes.docx'
        document.save(path)
        inventory(self.course, [path])
        source = read_json(self.course / '_工作区/资料索引.json')['sources'][0]
        extract(self.course, source['id'], options())
        unit = self.mapping(source)['units'][0]
        text = ''.join((self.course / c).read_text(encoding='utf-8') for c in unit['chunks'])
        self.assertIn('distance per unit time', text)
        self.assertIn('Distance\tTime', text)

    def test_assembly_rebases_attachment_links_without_rewriting_code(self):
        self.valid_package()
        fragment = self.course / '_工作区/章节知识/KC-001.md'
        atomic_text(self.course / '附件/说明.md', '# Evidence\n')
        with fragment.open('a', encoding='utf-8') as f:
            f.write('\n[Evidence](../../附件/说明.md)\n\n`[Inline code](../../example.md)`\n\n```text\n[Code](../../untouched.md)\n```\n')
        result = build_content(self.course)
        self.assertIn('[Evidence](%E9%99%84%E4%BB%B6/%E8%AF%B4%E6%98%8E.md)', result)
        self.assertIn('[Code](../../untouched.md)', result)
        self.assertIn('`[Inline code](../../example.md)`', result)
        atomic_text(self.course / '知识内容.md', result)
        self.assertTrue(validate(self.course)['passed'])

    def test_archived_originals_and_duplicate_detection(self):
        original = self.course / '原始资料/.notes.md'
        original.parent.mkdir(parents=True)
        original.write_text('Same material', encoding='utf-8')
        (self.inputs / 'copy.md').write_text('Same material', encoding='utf-8')
        inventory(self.course, [self.inputs, original.parent])
        sources = read_json(self.course / '_工作区/资料索引.json')['sources']
        self.assertEqual(len(sources), 2)
        self.assertEqual(sum('duplicate_of' in s for s in sources), 1)

    def test_valid_structure_and_lossless_assembly(self):
        self.valid_package()
        fragment = (self.course / '_工作区/章节知识/KC-001.md').read_text(encoding='utf-8')
        self.assertIn(fragment, build_content(self.course))
        report = validate(self.course)
        self.assertTrue(report['passed'], report['errors'])

    def test_validator_detects_lost_coverage_broken_anchor_and_stale_assembly(self):
        self.valid_package()
        self.jsonl('覆盖台账.jsonl', [])
        atomic_text(self.course / '开始学习.md', '# Start\n[Broken](学习文档/01-Rates.md#missing)\n')
        with (self.course / '知识内容.md').open('a', encoding='utf-8') as f:
            f.write('Silent divergence')
        report = validate(self.course)
        self.assertFalse(report['passed'])
        self.assertTrue(any('尚未登记覆盖' in e for e in report['errors']))
        self.assertTrue(any('锚点不存在' in e for e in report['errors']))
        self.assertTrue(any('分片不一致' in e for e in report['errors']))

    def test_validator_rejects_unreviewed_visual_and_dependency_cycle(self):
        source, knowledge = self.valid_package()
        map_file = extraction_path(self.course, source) / '定位映射.json'
        mapping = read_json(map_file)
        mapping['units'][0]['visual_review_required'] = True
        write_json(map_file, mapping)
        knowledge['prerequisites'] = [knowledge['id']]
        self.jsonl('知识索引.jsonl', [knowledge])
        report = validate(self.course)
        self.assertTrue(any('视觉内容尚未复核' in e for e in report['errors']))
        self.assertTrue(any('存在循环' in e for e in report['errors']))

    def test_incomplete_assembly_and_path_escape_rejected(self):
        self.valid_package()
        index_path = self.course / '_工作区/章节索引.json'
        index = read_json(index_path)
        index['chapters'][0]['status'] = 'pending'
        write_json(index_path, index)
        with self.assertRaisesRegex(ValueError, 'incomplete'):
            build_content(self.course)
        self.assertIn('未完成', build_content(self.course, allow_partial=True))
        index['chapters'][0]['status'] = 'complete'
        index['chapters'][0]['path'] = '../outside.md'
        write_json(index_path, index)
        with self.assertRaisesRegex(ValueError, 'leaves'):
            build_content(self.course)


if __name__ == '__main__':
    unittest.main(verbosity=2)
